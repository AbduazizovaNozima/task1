import logging
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import BinaryIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.domain.errors import InvalidSpreadsheetError
from app.domain.models import QAPair, QASheet, SkippedRow

logger = logging.getLogger(__name__)

QUESTION_COLUMN = "question"
ANSWER_COLUMN = "answer"
HEADER_ROW = 1


def read_qa_sheet(source: Path | BinaryIO) -> QASheet:
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            columns = _column_positions(next(rows, ()))
            return _parse_rows(rows, columns[QUESTION_COLUMN], columns[ANSWER_COLUMN])
        finally:
            workbook.close()
    except (BadZipFile, InvalidFileException, KeyError, ParseError) as exc:
        raise InvalidSpreadsheetError("Fayl .xlsx formatida emas yoki buzilgan.") from exc


def _column_positions(header: Sequence[object]) -> dict[str, int]:
    names = [_text(cell).lower() for cell in header]
    missing = [column for column in (QUESTION_COLUMN, ANSWER_COLUMN) if column not in names]
    if missing:
        raise InvalidSpreadsheetError(f"Majburiy ustun(lar) topilmadi: {', '.join(missing)}.")
    return {column: names.index(column) for column in (QUESTION_COLUMN, ANSWER_COLUMN)}


def _parse_rows(rows: Iterable[Sequence[object]], question_at: int, answer_at: int) -> QASheet:
    pairs: list[QAPair] = []
    skipped: list[SkippedRow] = []
    for number, row in enumerate(rows, start=HEADER_ROW + 1):
        question, answer = _cell(row, question_at), _cell(row, answer_at)
        if question and answer:
            pairs.append(QAPair(question=question, answer=answer))
            continue
        reason = "empty question" if not question else "empty answer"
        skipped.append(SkippedRow(row=number, reason=reason))
        logger.warning("Row %d skipped: %s", number, reason)
    return QASheet(pairs=pairs, skipped=skipped)


def _cell(row: Sequence[object], index: int) -> str:
    return _text(row[index]) if index < len(row) else ""


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()
